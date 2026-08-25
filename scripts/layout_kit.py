#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""layout.json(좌표 단일 진실 소스) → 스케치 HTML / 합성 HTML / 편집가능 PPTX.

사용법:
    python layout_kit.py sketch  <layout.json> [-o out.html]   # 회색 박스 구도 스케치 (codex -i 지시서)
    python layout_kit.py render  <layout.json> [-o out.html]   # 배경 + 실제 한글 텍스트 합성 (미리보기/최종 PNG 소스)
    python layout_kit.py pptx    <layout.json> [more.json ...] [-o out.pptx]  # 여러 장 배경 + 네이티브 텍스트박스
    python layout_kit.py capture <in.html> <out.png> --size WxH  # Chrome 헤드리스 캡처 + 1:1 검증
    python layout_kit.py demo                                   # 자가 점검 (임시 폴더, 흔적 없음)

layout.json 스키마:
    {"canvas": {"w": 1920, "h": 1080},
     "background": "gpt-image/xxx.png",          # sketch에는 불필요, render/pptx에 필요
     "labels": [{"id": "title", "text": "고객만족도 향상",   # \n 줄바꿈 지원
                 "x_pct": 6, "y_pct": 8, "w_pct": 46, "h_pct": 14,
                 "font": "Malgun Gothic", "size_pt": 40, "bold": true,
                 "color": "#111827", "align": "left", "scrim": false}]}

좌표 계약: 슬라이드 폭 13.333in 고정(16:9 기준), 높이는 캔버스 비율 추종.
  1pt = canvas.w/960 px  (13.333in × 72pt/in = 960pt ↔ canvas.w px)
  → render HTML의 px 폰트와 pptx의 pt 폰트가 같은 시각 비율이 된다.
# ponytail: layout.json 1개=슬라이드 1장. pptx는 같은 canvas 종횡비의 여러 파일을 입력 순서대로 한 덱에 묶는다.
"""
import argparse
import base64
import html as html_mod
import json
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zlib
from pathlib import Path

DEFAULT_FONT = "Malgun Gothic"
SLIDE_W_EMU = 12192000  # 표준 16:9 프리셋 폭. Inches(13.333)=12191695로 미세하게 어긋나므로 EMU 직접 사용
SCRIM_RGB = (0x11, 0x18, 0x27)  # 어두운 반투명 패널 (55%)
CHROME_CANDIDATES = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
]


# ---------- 공통 ----------

def load_layout(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    for key in ("canvas", "labels"):
        if key not in data:
            sys.exit(f"layout.json에 {key} 없음")
    for lb in data["labels"]:
        for k in ("text", "x_pct", "y_pct", "w_pct", "h_pct", "size_pt"):
            if k not in lb:
                sys.exit(f"label({lb.get('id', '?')})에 {k} 없음")
    return data


def png_size(path: Path) -> tuple:
    head = path.read_bytes()[:26]
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        sys.exit(f"PNG 아님: {path}")
    return struct.unpack(">II", head[16:24])


def find_chrome():
    for c in CHROME_CANDIDATES:
        if Path(c).exists():
            return c
    return shutil.which("chrome") or shutil.which("chromium") or shutil.which("msedge")


def pt_to_px(layout: dict) -> float:
    return layout["canvas"]["w"] / 960.0


def cover_crop(bw: int, bh: int, cw: int, ch: int) -> tuple:
    """배경이 캔버스 비율과 다를 때 center cover-crop 분율 (f_lr, f_tb).
    render(CSS object-fit:cover)와 pptx(crop_*)가 같은 수식을 공유한다."""
    ro, rt = bw / bh, cw / ch
    if abs(ro - rt) / rt <= 0.01:
        return 0.0, 0.0
    if ro > rt:   # 원본이 더 넓다 → 좌우 크롭
        return (1 - rt / ro) / 2, 0.0
    return 0.0, (1 - ro / rt) / 2


def label_box_px(layout: dict, lb: dict) -> tuple:
    cw, ch = layout["canvas"]["w"], layout["canvas"]["h"]
    return (round(cw * lb["x_pct"] / 100), round(ch * lb["y_pct"] / 100),
            round(cw * lb["w_pct"] / 100), round(ch * lb["h_pct"] / 100))


# ---------- sketch / render (stdlib) ----------

def html_page(layout: dict, body: str) -> str:
    cw, ch = layout["canvas"]["w"], layout["canvas"]["h"]
    return ("<!doctype html><meta charset='utf-8'>"
            f"<style>*{{margin:0}}body{{width:{cw}px;height:{ch}px;"
            "position:relative;overflow:hidden;background:#fff}</style>" + body)


def sketch_html(layout: dict) -> str:
    boxes = []
    for lb in layout["labels"]:
        x, y, w, h = label_box_px(layout, lb)
        boxes.append(f'<div style="position:absolute;left:{x}px;top:{y}px;'
                     f'width:{w}px;height:{h}px;background:#b3b3b3"></div>')
    return html_page(layout, "".join(boxes))


def render_html(layout: dict, base: Path) -> str:
    factor = pt_to_px(layout)
    # data URI 인라인 — file:// 이미지의 로딩 타이밍/경로 문제를 원천 제거 (부분 페인트 실측 사례)
    # object-fit:cover — 배경 비율이 캔버스와 달라도 stretch 왜곡 없이 center crop (pptx crop_*과 동일 수식)
    b64 = base64.b64encode(base.read_bytes()).decode("ascii")
    parts = [f'<img src="data:image/png;base64,{b64}" '
             'style="position:absolute;inset:0;width:100%;height:100%;'
             'object-fit:cover;object-position:center">']
    for lb in layout["labels"]:
        x, y, w, h = label_box_px(layout, lb)
        font = lb.get("font", DEFAULT_FONT)
        # 주의: style 속성은 큰따옴표, font-family는 작은따옴표 — 섞이면 속성이 중간에 끊긴다 (실측 버그)
        style = (f"position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;"
                 f"font-family:'{font}','{DEFAULT_FONT}',sans-serif;"
                 f"font-size:{lb['size_pt'] * factor:.1f}px;"
                 f"font-weight:{700 if lb.get('bold') else 400};"
                 f"color:{lb.get('color', '#111827')};"
                 f"text-align:{lb.get('align', 'left')};"
                 "line-height:1.2;white-space:pre-wrap")
        if lb.get("scrim"):
            style += f";background:rgba({SCRIM_RGB[0]},{SCRIM_RGB[1]},{SCRIM_RGB[2]},.55)"
        parts.append(f'<div style="{style}">{html_mod.escape(str(lb["text"]))}</div>')
    return html_page(layout, "".join(parts))


# ---------- pptx (python-pptx 필요 시에만 import) ----------

def _draw_pptx_slide(prs, layout: dict, slide_w: int, slide_h: int):
    """layout 하나를 blank slide 한 장에 그린다. 모든 장이 같은 렌더링 규약을 공유한다."""
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.oxml.ns import qn
    from lxml import etree

    cw, ch = layout["canvas"]["w"], layout["canvas"]["h"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    bg = layout.get("background")
    if bg:
        p = Path(bg)
        if not p.exists():
            sys.exit(f"배경 PNG 없음: {p}")
        pic = slide.shapes.add_picture(str(p), 0, 0, width=prs.slide_width, height=prs.slide_height)
        bw, bh = png_size(p)
        f_lr, f_tb = cover_crop(bw, bh, cw, ch)  # 비율 불일치 시 stretch 대신 center crop
        if f_lr:
            pic.crop_left = pic.crop_right = f_lr
        if f_tb:
            pic.crop_top = pic.crop_bottom = f_tb

    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    for lb in layout["labels"]:
        box = slide.shapes.add_textbox(
            Emu(round(slide_w * lb["x_pct"] / 100)), Emu(round(slide_h * lb["y_pct"] / 100)),
            Emu(round(slide_w * lb["w_pct"] / 100)), Emu(round(slide_h * lb["h_pct"] / 100)))
        if lb.get("scrim"):
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(*SCRIM_RGB)
            srgb = box._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(srgb, qn("a:alpha")).set("val", "55000")  # 55% 불투명
        tf = box.text_frame
        tf.word_wrap = True                     # 기본값 wrap="none" — 켜지 않으면 라벨이 폭 무시하고 한 줄로 뻗음
        tf.auto_size = MSO_AUTO_SIZE.NONE       # 기본값 spAutoFit — 끄지 않으면 PPT가 박스를 제멋대로 리사이즈
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        font_name = lb.get("font", DEFAULT_FONT)
        for i, line in enumerate(str(lb["text"]).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align_map.get(lb.get("align", "left"), PP_ALIGN.LEFT)
            p.line_spacing = 1.2                # CSS line-height:1.2와 패리티
            r = p.add_run()
            r.text = line
            r.font.name = font_name          # a:latin — 라틴 글리프용
            r.font.size = Pt(lb["size_pt"])
            r.font.bold = bool(lb.get("bold"))
            r.font.color.rgb = RGBColor.from_string(lb.get("color", "#111827").lstrip("#"))
            # 한글은 a:ea(East Asian) typeface를 따로 박아야 지정 폰트가 적용된다
            # (기본 테마의 ea가 빈 문자열이라 a:latin만으로는 뷰어 폴백 폰트로 렌더됨)
            rPr = r._r.get_or_add_rPr()
            rPr.set("lang", "ko-KR")
            latin = rPr.find(qn("a:latin"))
            ea = etree.SubElement(rPr, qn("a:ea"))
            ea.set("typeface", font_name)
            cs = etree.SubElement(rPr, qn("a:cs"))
            cs.set("typeface", font_name)
            if latin is not None:
                latin.addnext(ea)   # 스키마 순서: latin → ea → cs
                ea.addnext(cs)


def build_pptx(layout, out: Path, layout_names=None):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("python-pptx 없음 — pip install python-pptx 후 재실행 (또는 라벨 가이드 표 폴백)")

    layouts = [layout] if isinstance(layout, dict) else list(layout)
    if not layouts:
        sys.exit("pptx에는 layout.json이 하나 이상 필요")

    names = list(layout_names) if layout_names is not None else []
    first_cw = layouts[0]["canvas"]["w"]
    first_ch = layouts[0]["canvas"]["h"]
    first_name = str(names[0]) if names else "슬라이드 1"
    for index, current in enumerate(layouts[1:], start=2):
        cw = current["canvas"]["w"]
        ch = current["canvas"]["h"]
        if cw * first_ch != first_cw * ch:
            current_name = str(names[index - 1]) if index <= len(names) else f"슬라이드 {index}"
            sys.exit(
                "pptx canvas 종횡비 불일치: "
                f"{current_name} ({cw}x{ch}) != {first_name} ({first_cw}x{first_ch})")

    slide_w = SLIDE_W_EMU
    slide_h = round(SLIDE_W_EMU * first_ch / first_cw)
    prs = Presentation()
    prs.slide_width = Emu(slide_w)
    prs.slide_height = Emu(slide_h)
    for current in layouts:
        _draw_pptx_slide(prs, current, slide_w, slide_h)
    prs.save(str(out))


# ---------- capture (8de8e4d capture_overlay.py 기반 + 뷰포트 결손 보정) ----------
# 실측: 현 Chrome headless=new에서 --window-size는 창 장식 포함 크기다.
# 뷰포트 = window-size - (스크롤바 예약 16, 툴바 95) 같은 결손이 생기고, 스크린샷
# 비트맵은 window-size 그대로 나와 "크기 검증은 통과하지만 하단·우측이 빈" PNG가 된다.
# → 결손(dx,dy)을 프로브로 실측해 window-size를 키우고, 비트맵을 목표 크기로 크롭한다.

_DEFICIT_CACHE = {}


def viewport_deficit(chrome: str, w: int, h: int) -> tuple:
    """요청 창 크기 대비 실제 뷰포트 결손 (dx, dy)를 실측한다. 프로세스 내 메모이즈."""
    key = (chrome, w, h)
    if key in _DEFICIT_CACHE:
        return _DEFICIT_CACHE[key]
    probe = Path(tempfile.mkdtemp(prefix="lk_probe_")) / "probe.html"
    probe.write_text(
        "<!doctype html><body><script>document.body.setAttribute("
        "'data-vp',innerWidth+'x'+innerHeight)</script>", encoding="utf-8")
    r = subprocess.run(
        [chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars",
         "--force-device-scale-factor=1", f"--window-size={w},{h}",
         "--dump-dom", probe.as_uri()],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
    m = re.search(r'data-vp="(\d+)x(\d+)"', r.stdout)
    if not m:
        sys.exit("뷰포트 프로브 실패 — Chrome --dump-dom 출력에 data-vp 없음")
    dx, dy = w - int(m.group(1)), h - int(m.group(2))
    if dx < 0 or dy < 0 or dx > 100 or dy > 300:
        sys.exit(f"뷰포트 결손 비정상: dx={dx}, dy={dy}")
    _DEFICIT_CACHE[key] = (dx, dy)
    return dx, dy


def crop_png(src: Path, dst: Path, w: int, h: int):
    """디코드 → 좌상단 w×h 크롭 → filter-0 재인코딩 (stdlib)."""
    sw, sh, ch, buf = decode_png(src)
    if sw < w or sh < h:
        sys.exit(f"크롭 불가: 원본 {(sw, sh)} < 목표 {(w, h)}")
    stride = sw * ch
    rows = b"".join(b"\x00" + bytes(buf[y * stride:y * stride + w * ch]) for y in range(h))
    def chunk(tag, data):
        c = tag + data
        return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c))
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2 if ch == 3 else 6, 0, 0, 0)
    dst.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
                    + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))


def check_body_size(html_path: Path, w: int, h: int):
    text = html_path.read_text(encoding="utf-8")
    m = re.search(r"body\s*\{([^}]*)\}", text, re.S)
    if not m:
        sys.exit("HTML에 body{...} 규칙이 없음 — body 크기 = 캔버스 픽셀 규약")
    mw = re.search(r"width:\s*(\d+)px", m.group(1))
    mh = re.search(r"height:\s*(\d+)px", m.group(1))
    if not (mw and mh):
        sys.exit("body에 width/height px 선언이 없음")
    declared = (int(mw.group(1)), int(mh.group(1)))
    if declared != (w, h):
        sys.exit(f"body 선언 크기 {declared} ≠ 캔버스 {(w, h)} — 1:1 규칙 위반(잘림/여백 발생)")


def capture(html_path: Path, out: Path, w: int, h: int) -> dict:
    html_path, out = html_path.resolve(), out.resolve()  # Chrome은 상대경로를 자기 cwd 기준으로 푼다
    if not html_path.exists():
        sys.exit(f"HTML 없음: {html_path}")
    check_body_size(html_path, w, h)
    chrome = find_chrome()
    if not chrome:
        sys.exit("Chrome을 찾지 못함 — SKILL.md 오류표의 폴백을 사용하세요")
    dx, dy = viewport_deficit(chrome, w, h)
    raw = out.with_suffix(".raw.png")
    r = subprocess.run(
        [chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars",
         "--force-device-scale-factor=1", f"--screenshot={raw}",
         f"--window-size={w + dx},{h + dy}", "--virtual-time-budget=3000",
         html_path.as_uri()],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
    if not raw.exists():
        sys.exit(f"캡처 실패: {r.stderr.strip()[-500:]}")
    size = png_size(raw)
    if size == (w, h):          # 결손 없는 환경(리눅스 등)
        raw.replace(out)
    elif size == (w + dx, h + dy):
        crop_png(raw, out, w, h)
        raw.unlink()
    else:
        sys.exit(f"예상 외 비트맵 크기: {size} (기대 {(w, h)} 또는 {(w + dx, h + dy)})")
    return {"out": str(out), "size": [w, h], "viewport_deficit": [dx, dy]}


# ---------- demo ----------

def decode_png(path: Path):
    """8-bit RGB/RGBA 비인터레이스 PNG 디코더 — demo 픽셀 검증 전용 (Pillow 회피)."""
    data = path.read_bytes()
    pos, idat, w, h, ch = 8, b"", 0, 0, 3
    while pos < len(data):
        ln, typ = struct.unpack(">I4s", data[pos:pos + 8])
        chunk = data[pos + 8:pos + 8 + ln]
        if typ == b"IHDR":
            w, h, depth, ctype = struct.unpack(">IIBB", chunk[:10])
            if depth != 8 or ctype not in (2, 6) or chunk[12] != 0:
                sys.exit("지원 외 PNG 형식 (8-bit RGB/RGBA 비인터레이스만)")
            ch = 3 if ctype == 2 else 4
        elif typ == b"IDAT":
            idat += chunk
        elif typ == b"IEND":
            break
        pos += 12 + ln
    raw = zlib.decompress(idat)
    stride = w * ch
    out = bytearray(h * stride)
    prev = bytearray(stride)
    for row in range(h):
        f = raw[row * (stride + 1)]
        line = bytearray(raw[row * (stride + 1) + 1:(row + 1) * (stride + 1)])
        for i in range(stride):
            a = line[i - ch] if i >= ch else 0
            b = prev[i]
            c = prev[i - ch] if i >= ch else 0
            if f == 1:
                line[i] = (line[i] + a) & 0xFF
            elif f == 2:
                line[i] = (line[i] + b) & 0xFF
            elif f == 3:
                line[i] = (line[i] + (a + b) // 2) & 0xFF
            elif f == 4:
                p = a + b - c
                pa, pb, pc = abs(p - a), abs(p - b), abs(p - c)
                line[i] = (line[i] + (a if pa <= pb and pa <= pc else b if pb <= pc else c)) & 0xFF
        out[row * stride:(row + 1) * stride] = line
        prev = line
    return w, h, ch, out


def pixel(decoded, x: int, y: int) -> tuple:
    w, _h, ch, buf = decoded
    o = (y * w + x) * ch
    return tuple(buf[o:o + 3])


def near(a: tuple, b: tuple, tol: int = 12) -> bool:
    return all(abs(x - y) <= tol for x, y in zip(a, b))


def write_solid_png(path: Path, w: int, h: int, rgb: tuple):
    row = b"\x00" + bytes(rgb) * w
    def chunk(tag, data):
        c = tag + data
        return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c))
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
                     + chunk(b"IDAT", zlib.compress(row * h)) + chunk(b"IEND", b""))


def demo() -> dict:
    tmp = Path(tempfile.mkdtemp(prefix="layout_kit_demo_"))
    base = tmp / "base.png"
    write_solid_png(base, 640, 360, (221, 231, 227))
    layout = {
        "canvas": {"w": 640, "h": 360},
        "background": str(base),
        "labels": [
            {"id": "title", "text": "고객만족도 향상", "x_pct": 6, "y_pct": 8,
             "w_pct": 60, "h_pct": 18, "size_pt": 20, "bold": True,
             "color": "#111827", "align": "left"},
            {"id": "sub", "text": "2026년 상반기\n실적 보고", "x_pct": 6, "y_pct": 70,
             "w_pct": 40, "h_pct": 22, "size_pt": 11, "color": "#FFFFFF",
             "align": "left", "scrim": True},
        ],
    }
    lp = tmp / "layout.json"
    lp.write_text(json.dumps(layout, ensure_ascii=False), encoding="utf-8")

    sk = sketch_html(layout)
    assert sk.count("position:absolute") == 2 and "#b3b3b3" in sk
    (tmp / "sketch.html").write_text(sk, encoding="utf-8")
    rd = render_html(layout, base)
    assert "고객만족도 향상" in rd and "rgba(17,24,39,.55)" in rd
    assert f"{20 * 640 / 960:.1f}px" in rd  # pt→px 계약: 20pt = 13.3px @640
    (tmp / "render.html").write_text(rd, encoding="utf-8")

    checks = {"sketch": "ok", "render": "ok"}

    try:
        import pptx  # noqa: F401
        out = tmp / "out.pptx"
        build_pptx(layout, out)
        from pptx import Presentation
        prs = Presentation(str(out))
        assert prs.slide_height == round(SLIDE_W_EMU * 360 / 640)  # 6858000 = 표준 16:9 높이
        shapes = list(prs.slides[0].shapes)
        assert len(shapes) == 3  # 배경 그림 + 텍스트박스 2
        tb = [s for s in shapes if s.has_text_frame]
        assert tb[0].text_frame.text == "고객만족도 향상"
        assert tb[0].left == round(SLIDE_W_EMU * 0.06)
        assert prs.slide_width == 12192000  # 표준 16:9 EMU 정확 일치
        import zipfile
        xml = zipfile.ZipFile(out).read("ppt/slides/slide1.xml").decode("utf-8")
        assert '<a:ea typeface="Malgun Gothic"' in xml  # 한글 폰트 적용 핵심
        assert '<a:cs typeface="Malgun Gothic"' in xml
        assert '<a:alpha val="55000"' in xml            # scrim 반투명
        assert "spAutoFit" not in xml                   # auto_size=NONE — 박스 크기 계약 보존
        assert 'wrap="square"' in xml                   # word_wrap=True
        # cover-crop 수식 단위 검증 (1536×1024 배경 → 16:9 캔버스: 상하 7.8125% 크롭)
        assert cover_crop(1536, 1024, 1920, 1080) == (0.0, (1 - (1536 / 1024) / (1920 / 1080)) / 2)
        assert cover_crop(1920, 1080, 640, 360) == (0.0, 0.0)  # 동일 비율 → 크롭 없음

        multi_paths = []
        expected_texts = []
        for index in range(1, 4):
            marker = f"멀티 슬라이드 {index}"
            slide_layout = {
                **layout,
                "canvas": {"w": 640 * index, "h": 360 * index},
                "labels": [dict(lb) for lb in layout["labels"]],
            }
            slide_layout["labels"][0]["text"] = marker
            slide_path = tmp / f"slide{index}.json"
            slide_path.write_text(json.dumps(slide_layout, ensure_ascii=False), encoding="utf-8")
            multi_paths.append(slide_path)
            expected_texts.append(marker)

        multi_out = tmp / "multi.pptx"
        multi_run = subprocess.run(
            [sys.executable, str(Path(__file__).resolve()), "pptx",
             *(str(path) for path in multi_paths), "-o", str(multi_out)],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
        assert multi_run.returncode == 0, f"멀티 pptx 생성 실패: {multi_run.stderr}"
        multi_prs = Presentation(str(multi_out))
        assert len(multi_prs.slides) == 3, f"멀티 pptx 슬라이드 수: {len(multi_prs.slides)}"
        multi_zip = zipfile.ZipFile(multi_out)
        for index, slide in enumerate(multi_prs.slides):
            assert len(slide.shapes) == 3  # 장마다 배경 그림 + 텍스트박스 2
            slide_text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            assert expected_texts[index] in slide_text
            assert all(marker not in slide_text for marker in expected_texts if marker != expected_texts[index])
            slide_xml = multi_zip.read(f"ppt/slides/slide{index + 1}.xml").decode("utf-8")
            assert '<a:ea typeface="Malgun Gothic"' in slide_xml
            assert '<a:alpha val="55000"' in slide_xml
            assert "spAutoFit" not in slide_xml
            assert 'wrap="square"' in slide_xml

        bad_layout = {**layout, "canvas": {"w": 640, "h": 400}}
        bad_path = tmp / "slide-bad-ratio.json"
        bad_path.write_text(json.dumps(bad_layout, ensure_ascii=False), encoding="utf-8")
        bad_out = tmp / "bad-ratio.pptx"
        bad_run = subprocess.run(
            [sys.executable, str(Path(__file__).resolve()), "pptx",
             str(multi_paths[0]), str(bad_path), "-o", str(bad_out)],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
        assert bad_run.returncode != 0, "canvas 종횡비 불일치가 중단되지 않음"
        assert "pptx canvas 종횡비 불일치" in bad_run.stderr, bad_run.stderr
        assert not bad_out.exists(), "종횡비 불일치인데 pptx가 생성됨"

        checks["pptx"] = "ok"
        checks["pptx_multi"] = "ok"
        checks["pptx_ratio_guard"] = "ok"
    except ImportError:
        checks["pptx"] = "skipped (python-pptx 없음)"

    if find_chrome():
        capture(tmp / "sketch.html", tmp / "sketch.png", 640, 360)
        capture(tmp / "render.html", tmp / "render.png", 640, 360)
        # 픽셀 프로브 — 눈대중이 아니라 좌표별 색으로 검증 (부분 페인트·CSS 유실 실측 재발 방지)
        skd = decode_png(tmp / "sketch.png")
        assert near(pixel(skd, 100, 300), (179, 179, 179)), "sketch 부제 박스 자리(#b3b3b3) 불일치"
        assert near(pixel(skd, 480, 200), (255, 255, 255)), "sketch 빈 영역이 흰색 아님"
        rdd = decode_png(tmp / "render.png")
        assert near(pixel(rdd, 620, 350), (221, 231, 227)), "render 우하단이 배경색 아님 — 배경 부분 페인트 의심"
        scrim_expect = tuple(round(0.55 * s + 0.45 * b) for s, b in zip(SCRIM_RGB, (221, 231, 227)))
        assert near(pixel(rdd, 280, 326), scrim_expect), "render scrim 반투명 합성색 불일치"
        checks["capture"] = "ok"
        checks["pixel_probe"] = "ok"
    else:
        checks["capture"] = "skipped (Chrome 없음)"

    checks["dir"] = str(tmp)
    return checks


# ---------- CLI ----------

def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("sketch", "render"):
        p = sub.add_parser(name)
        p.add_argument("layout", type=Path)
        p.add_argument("-o", "--out", type=Path)
    pp = sub.add_parser("pptx")
    pp.add_argument("layouts", nargs="+", type=Path)
    pp.add_argument("-o", "--out", type=Path)
    pc = sub.add_parser("capture")
    pc.add_argument("html", type=Path)
    pc.add_argument("out", type=Path)
    pc.add_argument("--size", required=True, help="WxH 예: 1920x1080")
    sub.add_parser("demo")
    a = ap.parse_args()

    if a.cmd == "demo":
        print(json.dumps(demo(), ensure_ascii=False))
        return
    if a.cmd == "capture":
        w, h = (int(v) for v in a.size.lower().split("x"))
        print(json.dumps(capture(a.html, a.out, w, h), ensure_ascii=False))
        return
    if a.cmd == "pptx":
        layouts = [load_layout(path) for path in a.layouts]
        stem = a.layouts[0].with_suffix("")
        out = a.out or Path(f"{stem}.pptx")
        build_pptx(layouts[0] if len(layouts) == 1 else layouts, out, a.layouts)
        print(json.dumps({"out": str(out)}, ensure_ascii=False))
        return

    layout = load_layout(a.layout)
    stem = a.layout.with_suffix("")
    if a.cmd == "sketch":
        out = a.out or Path(f"{stem}-sketch.html")
        out.write_text(sketch_html(layout), encoding="utf-8")
    elif a.cmd == "render":
        bg = layout.get("background")
        if not bg or not Path(bg).exists():
            sys.exit(f"배경 PNG 없음: {bg} — render에는 background가 필요")
        out = a.out or Path(f"{stem}-render.html")
        out.write_text(render_html(layout, Path(bg)), encoding="utf-8")
    print(json.dumps({"out": str(out)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
